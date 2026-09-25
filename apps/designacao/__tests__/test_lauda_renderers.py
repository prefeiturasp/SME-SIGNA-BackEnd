"""Testes dos renderizadores da lauda (PDF via ReportLab, Word via python-docx).

Trabalham sobre um ``DocumentoLauda`` fixo, sem banco: verificam que o
conteúdo chega ao arquivo, que o PDF escapa markup e que o Word preserva
o texto literalmente.
"""

from io import BytesIO

import pytest
from docx import Document
from reportlab.lib.colors import toColor
from reportlab.platypus import Paragraph

from apps.designacao.services.lauda_renderers import (
    DocumentoLauda,
    SecaoLauda,
    Trecho,
    montar_story,
    renderizar_docx,
    renderizar_pdf,
)

TEXTO_COM_MARKUP = "Cargo <b>X</b> & <img src='file:///etc/passwd'/>"


@pytest.fixture
def documento() -> DocumentoLauda:
    return DocumentoLauda(
        aviso="Aviso de teste",
        titulo="Lauda de teste",
        sei="6018.2026/0001234-5",
        subtitulo="Gerada em 18/09/2026 às 10:00 — 2 ato(s)",
        secoes=[
            SecaoLauda(
                titulo="PORTARIA Nº 5/2026 – DESIGNAÇÃO",
                sei="6018.2026/0000005-0",
                blocos=[
                    [
                        [Trecho("Primeira linha.")],
                        [Trecho("Segunda linha.")],
                    ],
                    [
                        [
                            Trecho("Fim "),
                            Trecho("negrito", negrito=True),
                            Trecho(" "),
                            Trecho("itálico", italico=True),
                            Trecho(" "),
                            Trecho("ambos", negrito=True, italico=True),
                            Trecho(" "),
                            Trecho("sub", sublinhado=True),
                            Trecho("."),
                        ]
                    ],
                ],
            ),
            SecaoLauda(
                titulo="APOSTILA",
                sei="6018.2026/0000099-0",
                blocos=[[[Trecho(TEXTO_COM_MARKUP)]]],
            ),
        ],
    )


def _paragrafos(flowables) -> list[Paragraph]:
    return [item for item in flowables if isinstance(item, Paragraph)]


def _textos(flowables) -> list[str]:
    return [p.text for p in _paragrafos(flowables)]


class TestPdf:
    """Renderização em PDF."""

    def test_story_comeca_com_aviso_titulo_sei_e_subtitulo(self, documento):
        textos = _textos(montar_story(documento))
        assert textos[:4] == [
            "Aviso de teste",
            "Lauda de teste",
            "SEI nº 6018.2026/0001234-5",
            "Gerada em 18/09/2026 às 10:00 — 2 ato(s)",
        ]

    def test_story_aviso_vermelho_e_negrito(self, documento):
        aviso = _paragrafos(montar_story(documento))[0]
        assert toColor(aviso.style.textColor) == toColor("#C00000")
        assert aviso.style.fontName == "Helvetica-Bold"

    def test_story_sem_aviso_quando_none(self, documento):
        from dataclasses import replace

        textos = _textos(montar_story(replace(documento, aviso=None)))
        assert textos[0] == "Lauda de teste"

    def test_story_nao_imprime_cabecalho_por_ato(self, documento):
        textos = _textos(montar_story(documento))
        assert "PORTARIA Nº 5/2026 – DESIGNAÇÃO" not in textos
        assert "APOSTILA" not in textos
        assert "SEI 6018.2026/0000005-0" not in textos

    def test_story_linhas_viram_br(self, documento):
        textos = _textos(montar_story(documento))
        assert "Primeira linha.<br/>Segunda linha." in textos

    def test_story_escapa_markup(self, documento):
        # O parser de Paragraph do ReportLab interpreta tags — texto do
        # banco nunca pode chegar cru.
        textos = _textos(montar_story(documento))
        assert not any("<img" in t for t in textos)
        assert any("&lt;b&gt;X&lt;/b&gt; &amp;" in t for t in textos)

    def test_story_aplica_formatacao_com_tags_do_reportlab(self, documento):
        textos = _textos(montar_story(documento))
        assert (
            "Fim <b>negrito</b> <i>itálico</i> <i><b>ambos</b></i> "
            "<u>sub</u>."
        ) in textos

    def test_gera_pdf_valido(self, documento):
        pdf = renderizar_pdf(documento)
        assert pdf.startswith(b"%PDF-")
        assert b"%%EOF" in pdf[-64:]


class TestDocx:
    """Renderização em Word."""

    @pytest.fixture
    def docx(self, documento):
        return Document(BytesIO(renderizar_docx(documento)))

    def test_comeca_com_aviso_titulo_sei_e_subtitulo(self, docx):
        textos = [p.text for p in docx.paragraphs]
        assert textos[:4] == [
            "Aviso de teste",
            "Lauda de teste",
            "SEI nº 6018.2026/0001234-5",
            "Gerada em 18/09/2026 às 10:00 — 2 ato(s)",
        ]

    def test_aviso_vermelho_e_negrito(self, docx):
        (run,) = docx.paragraphs[0].runs
        assert run.bold
        assert str(run.font.color.rgb) == "C00000"

    def test_sem_aviso_quando_none(self, documento):
        from dataclasses import replace

        docx = Document(
            BytesIO(renderizar_docx(replace(documento, aviso=None)))
        )
        assert docx.paragraphs[0].text == "Lauda de teste"

    def test_textos_dos_atos_na_ordem_sem_cabecalho(self, docx):
        textos = [p.text for p in docx.paragraphs]
        assert "PORTARIA Nº 5/2026 – DESIGNAÇÃO" not in textos
        assert "APOSTILA" not in textos
        i_primeiro = textos.index("Primeira linha.\nSegunda linha.")
        i_apostila = textos.index(TEXTO_COM_MARKUP)
        assert i_primeiro < i_apostila

    def test_linhas_do_paragrafo_viram_quebra_de_linha(self, docx):
        textos = [p.text for p in docx.paragraphs]
        assert "Primeira linha.\nSegunda linha." in textos

    def test_formatacao_dos_runs(self, docx):
        paragrafo = next(
            p for p in docx.paragraphs if p.text.startswith("Fim ")
        )
        runs = {r.text: r for r in paragrafo.runs}
        assert runs["negrito"].bold and not runs["negrito"].italic
        assert runs["itálico"].italic and not runs["itálico"].bold
        assert runs["ambos"].bold and runs["ambos"].italic
        assert runs["sub"].underline
        assert not runs["Fim "].bold and not runs["Fim "].italic

    def test_texto_preservado_literalmente(self, docx):
        # python-docx grava texto, não markup — nada a escapar.
        assert TEXTO_COM_MARKUP in [p.text for p in docx.paragraphs]

    def test_rodape_tem_campo_de_pagina(self, docx):
        rodape = docx.sections[0].footer.paragraphs[0]
        assert rodape.text.startswith("Página ")
        assert 'w:instr="PAGE"' in rodape._p.xml

    def test_metadados(self, docx):
        assert docx.core_properties.title == "Lauda de teste"
