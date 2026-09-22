"""Renderizadores da lauda de publicação no D.O. (PDF e Word).

Recebem um ``DocumentoLauda`` já montado pelo ``LaudaService`` e devolvem
os bytes do arquivo. Nenhuma regra de negócio vive aqui — só layout.
"""

from dataclasses import dataclass
from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph as DocxParagraph
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import (
    Flowable,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)

CONTENT_TYPE_PDF = "application/pdf"
CONTENT_TYPE_DOCX = (
    "application/vnd.openxmlformats-officedocument."
    "wordprocessingml.document"
)

_CINZA = "#555555"
_VERMELHO = "#C00000"
_MARGEM_PDF = 20 * mm


@dataclass(frozen=True)
class Trecho:
    """Um trecho de texto com formatação uniforme (um *run*)."""

    texto: str
    negrito: bool = False
    italico: bool = False
    sublinhado: bool = False


# Uma linha é uma sequência de trechos; um parágrafo, uma sequência de
# linhas (quebras simples preservadas).
Linha = list[Trecho]
Paragrafo = list[Linha]


@dataclass(frozen=True)
class SecaoLauda:
    """Um ato dentro da lauda.

    ``titulo`` e ``sei`` não são impressos por enquanto: a lauda vai ser
    dividida por DRE e depois por tipo de ato, e o cabeçalho de cada
    grupo ainda será definido. Ficam aqui porque o agrupamento vai
    precisar deles.

    Attributes:
        titulo: Cabeçalho do ato (``PORTARIA Nº 12/2026 – DESIGNAÇÃO``).
        sei: Número do processo SEI do ato.
        blocos: Parágrafos do texto SEI, já como trechos formatados.

    """

    titulo: str
    sei: str
    blocos: list[Paragrafo]


@dataclass(frozen=True)
class DocumentoLauda:
    """Conteúdo completo da lauda, independente de formato.

    Attributes:
        aviso: Alerta impresso antes de tudo, em vermelho e negrito
            (``None`` para omitir).
        titulo: Título da lauda.
        sei: Número SEI da lauda, impresso logo abaixo do título.
        subtitulo: Linha discreta de metadados (data de geração).
        secoes: Um item por ato, na ordem de impressão.

    """

    aviso: str | None
    titulo: str
    sei: str
    subtitulo: str
    secoes: list[SecaoLauda]


# ─── PDF ─────────────────────────────────────────────────────────────────────


def renderizar_pdf(documento: DocumentoLauda) -> bytes:
    """Gera o PDF da lauda com ReportLab.

    Args:
        documento: Conteúdo da lauda.

    Returns:
        bytes: Arquivo PDF.

    """
    buffer = BytesIO()
    pdf = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=_MARGEM_PDF,
        rightMargin=_MARGEM_PDF,
        topMargin=_MARGEM_PDF,
        bottomMargin=_MARGEM_PDF,
        title=documento.titulo,
        author="SIGNA",
    )
    pdf.build(
        montar_story(documento),
        onFirstPage=_rodape_pdf,
        onLaterPages=_rodape_pdf,
    )
    return buffer.getvalue()


def montar_story(documento: DocumentoLauda) -> list[Flowable]:
    """Monta a sequência de elementos do PDF.

    Todo texto passa por ``escape`` antes de virar ``Paragraph`` — o
    parser de markup do ReportLab já teve RCEs por conteúdo não
    confiável, e o texto SEI vem do banco.

    Args:
        documento: Conteúdo da lauda.

    Returns:
        list[Flowable]: Aviso, título, SEI, subtítulo e os textos dos
        atos em sequência.

    """
    estilos = _estilos_pdf()
    story: list[Flowable] = []
    if documento.aviso:
        story.append(Paragraph(escape(documento.aviso), estilos["aviso"]))
    story += [
        Paragraph(escape(documento.titulo), estilos["titulo"]),
        Paragraph(f"SEI nº {escape(documento.sei)}", estilos["sei"]),
        Paragraph(escape(documento.subtitulo), estilos["meta"]),
        Spacer(1, 8 * mm),
    ]

    for secao in documento.secoes:
        story.extend(
            Paragraph(
                "<br/>".join(_linha_pdf(linha) for linha in paragrafo),
                estilos["corpo"],
            )
            for paragrafo in secao.blocos
        )
        story.append(Spacer(1, 6 * mm))

    return story


def _linha_pdf(linha: Linha) -> str:
    """Monta o markup do ``Paragraph`` para uma linha (texto escapado)."""
    partes = []
    for trecho in linha:
        texto = escape(trecho.texto)
        if trecho.negrito:
            texto = f"<b>{texto}</b>"
        if trecho.italico:
            texto = f"<i>{texto}</i>"
        if trecho.sublinhado:
            texto = f"<u>{texto}</u>"
        partes.append(texto)
    return "".join(partes)


def _rodape_pdf(canvas: Canvas, documento: SimpleDocTemplate) -> None:
    largura, _ = documento.pagesize
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.drawRightString(
        largura - _MARGEM_PDF, 10 * mm, f"Página {documento.page}"
    )
    canvas.restoreState()


def _estilos_pdf() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "aviso": ParagraphStyle(
            "LaudaAviso",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=10,
            leading=13,
            textColor=_VERMELHO,
            alignment=TA_CENTER,
            spaceAfter=6 * mm,
        ),
        "titulo": ParagraphStyle(
            "LaudaTitulo",
            parent=base["Title"],
            fontSize=14,
            leading=18,
            alignment=TA_CENTER,
            spaceAfter=2 * mm,
        ),
        "sei": ParagraphStyle(
            "LaudaSei",
            parent=base["Normal"],
            fontSize=11,
            leading=14,
            alignment=TA_CENTER,
            spaceAfter=1 * mm,
        ),
        "meta": ParagraphStyle(
            "LaudaMeta",
            parent=base["Normal"],
            fontSize=9,
            textColor=_CINZA,
            alignment=TA_CENTER,
        ),
        "corpo": ParagraphStyle(
            "LaudaCorpo",
            parent=base["Normal"],
            fontSize=10,
            leading=14,
            alignment=TA_JUSTIFY,
            spaceAfter=3 * mm,
        ),
    }


# ─── Word ────────────────────────────────────────────────────────────────────


def renderizar_docx(documento: DocumentoLauda) -> bytes:
    """Gera o .docx da lauda com python-docx.

    Diferente do ReportLab, ``add_run`` grava o texto literalmente no
    XML (o lxml escapa) — não há markup a neutralizar.

    Args:
        documento: Conteúdo da lauda.

    Returns:
        bytes: Arquivo Word (.docx).

    """
    docx = Document()
    docx.core_properties.title = documento.titulo
    docx.core_properties.author = "SIGNA"

    normal = docx.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)

    _cabecalho_docx(docx, documento)
    for secao in documento.secoes:
        for bloco in secao.blocos:
            _paragrafo_docx(docx, bloco)
    _numerar_paginas_docx(docx)

    buffer = BytesIO()
    docx.save(buffer)
    return buffer.getvalue()


def _cabecalho_docx(docx: DocxDocument, documento: DocumentoLauda) -> None:
    """Aviso (opcional), título, SEI e subtítulo, todos centralizados."""
    if documento.aviso:
        aviso = docx.add_paragraph()
        aviso.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = aviso.add_run(documento.aviso)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(_VERMELHO.lstrip("#"))

    titulo = docx.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo.add_run(documento.titulo)
    run.bold = True
    run.font.size = Pt(14)

    sei = docx.add_paragraph()
    sei.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sei.add_run(f"SEI nº {documento.sei}").font.size = Pt(11)

    subtitulo = docx.add_paragraph()
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run_cinza(subtitulo, documento.subtitulo)
    docx.add_paragraph()


def _paragrafo_docx(docx: DocxDocument, bloco: Paragrafo) -> None:
    """Parágrafo justificado: quebra entre linhas, um run por trecho."""
    paragrafo = docx.add_paragraph()
    paragrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragrafo.paragraph_format.space_after = Pt(6)
    for indice, linha in enumerate(bloco):
        if indice:
            paragrafo.add_run().add_break()
        for trecho in linha:
            run = paragrafo.add_run(trecho.texto)
            run.bold = trecho.negrito or None
            run.italic = trecho.italico or None
            run.underline = trecho.sublinhado or None


def _run_cinza(paragrafo: DocxParagraph, texto: str) -> None:
    run = paragrafo.add_run(texto)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(_CINZA.lstrip("#"))


def _numerar_paginas_docx(docx: DocxDocument) -> None:
    """Insere ``Página N`` no rodapé via campo ``PAGE`` do Word."""
    rodape = docx.sections[0].footer.paragraphs[0]
    rodape.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rodape.add_run("Página ").font.size = Pt(8)

    campo = OxmlElement("w:fldSimple")
    campo.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    texto = OxmlElement("w:t")
    texto.text = "1"
    run.append(texto)
    campo.append(run)
    rodape._p.append(campo)
